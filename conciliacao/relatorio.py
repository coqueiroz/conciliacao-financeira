"""Geração do relatório de conciliação em Excel."""
from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

ORDEM_SITUACOES = [
    "Conciliado",
    "Valor divergente",
    "Pagamento duplicado",
    "Venda sem pagamento",
    "Pagamento sem venda",
    "Venda cancelada com pagamento",
]

COLUNAS = {
    "codigo_venda": "Código da venda",
    "data_venda": "Data da venda",
    "cliente": "Cliente",
    "valor_venda": "Valor da venda (R$)",
    "valor_pago": "Valor recebido (R$)",
    "diferenca": "Diferença (R$)",
    "qtd_pagamentos": "Qtd. pagamentos",
    "data_pagamento": "Data do pagamento",
    "documentos": "Documento(s)",
    "historico": "Histórico do extrato",
    "dias_em_aberto": "Dias em aberto",
}

AZUL = PatternFill("solid", start_color="1F3A5F")
FORMATO_MOEDA = '"R$" #,##0.00;[Red]-"R$" #,##0.00'


def _resumo(resultado: pd.DataFrame) -> pd.DataFrame:
    linhas = []
    for situacao in ORDEM_SITUACOES:
        grupo = resultado[resultado["situacao"] == situacao]
        linhas.append({
            "Situação": situacao,
            "Quantidade": len(grupo),
            "Valor das vendas (R$)": grupo["valor_venda"].sum(),
            "Valor recebido (R$)": grupo["valor_pago"].sum(),
            "Diferença (R$)": grupo["diferenca"].sum(),
        })
    return pd.DataFrame(linhas).round(2)


COLUNAS_DATA = {"Data da venda", "Data do pagamento"}


def _formatar(planilha, colunas_moeda: set[str]) -> None:
    for celula in planilha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = AZUL
        celula.alignment = Alignment(vertical="center", wrap_text=True)
    planilha.row_dimensions[1].height = 30
    planilha.freeze_panes = "A2"
    for i, coluna in enumerate(planilha.iter_cols(min_row=1, max_row=planilha.max_row), start=1):
        titulo = str(coluna[0].value)
        largura = max(len(str(c.value)) if c.value is not None else 0 for c in coluna)
        planilha.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 12), 45)
        if titulo in colunas_moeda:
            for celula in coluna[1:]:
                celula.number_format = FORMATO_MOEDA
        elif titulo in COLUNAS_DATA:
            planilha.column_dimensions[get_column_letter(i)].width = 14
            for celula in coluna[1:]:
                celula.number_format = "DD/MM/YYYY"


def gerar_excel(resultado: pd.DataFrame, invalidos: pd.DataFrame, caminho: Path, data_referencia: str) -> None:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    resumo = _resumo(resultado)
    moeda = {"Valor das vendas (R$)", "Valor recebido (R$)", "Diferença (R$)", "Valor da venda (R$)"}

    with pd.ExcelWriter(caminho, engine="openpyxl") as escritor:
        resumo.to_excel(escritor, sheet_name="Resumo", index=False, startrow=3)
        aba = escritor.sheets["Resumo"]
        aba["A1"] = f"Conciliação financeira — referência {pd.Timestamp(data_referencia):%d/%m/%Y}"
        aba["A1"].font = Font(bold=True, size=14)
        conciliados = int(resumo.loc[resumo["Situação"] == "Conciliado", "Quantidade"].iloc[0])
        faturadas = resultado["valor_venda"].notna() & (resultado["situacao"] != "Venda cancelada com pagamento")
        total_vendas = int(faturadas.sum())
        aba["A2"] = (f"{conciliados} de {total_vendas} vendas faturadas conciliadas "
                     f"({conciliados / total_vendas:.0%}). Abas seguintes: detalhes de cada pendência.")

        for situacao in ORDEM_SITUACOES[1:]:
            detalhe = resultado[resultado["situacao"] == situacao]
            colunas = [c for c in COLUNAS if detalhe[c].notna().any()]
            nome_aba = situacao[:31]
            detalhe[colunas].rename(columns=COLUNAS).to_excel(escritor, sheet_name=nome_aba, index=False)

        resultado[resultado["situacao"] == "Conciliado"][
            ["codigo_venda", "data_venda", "cliente", "valor_venda", "valor_pago", "data_pagamento", "documentos"]
        ].rename(columns=COLUNAS).to_excel(escritor, sheet_name="Conciliados", index=False)

        if not invalidos.empty:
            invalidos[["origem", "linha_arquivo", "motivo"]].rename(columns={
                "origem": "Arquivo", "linha_arquivo": "Linha", "motivo": "Problema"
            }).to_excel(escritor, sheet_name="Dados inválidos", index=False)

        for nome, planilha in escritor.sheets.items():
            if nome == "Resumo":
                continue
            _formatar(planilha, moeda)

        # a aba Resumo tem o cabeçalho na linha 4
        for celula in aba[4]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = AZUL
        for coluna in "BCDE":
            aba.column_dimensions[coluna].width = 22
        aba.column_dimensions["A"].width = 32
        for linha in aba.iter_rows(min_row=5, min_col=3, max_col=5):
            for celula in linha:
                celula.number_format = FORMATO_MOEDA

    log.info("Relatório salvo em %s", caminho)
